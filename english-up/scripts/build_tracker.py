#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""English up · 每日打卡表生成器

生成一个一年期英语学习打卡工作簿（4 个工作表）：
  打卡表    逐日登记：日期 / 星期 / 周次 / 阶段 / 五个板块分钟数 / 合计 / 目标 / 完成率 / 打卡 / 备注
  周汇总    按周自动 SUMIF 汇总计划与实际时长、完成率、打卡天数
  指标看板  8 项核心 KPI 的起点 / 目标 / 当前值
  资源清单  材料与工具，随进度勾选

用法：
  python build_tracker.py --out 每日打卡表.xlsx --start 2026-10-19 --weeks 52 \
      --prep-days 4 --app 不背单词 --exam-name CET-6 --daily-min 150

只依赖 openpyxl。生成后建议用 LibreOffice 或 Excel 打开一次让公式求值；
也可用 LibreOffice Kit：recalculate --input <out> --output <recalculated>。
"""
from __future__ import annotations

import argparse
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# 4 阶段 × 每周主题（通用版，可按需替换）
THEMES = [
    ("P1 地基重建", "发音与语法骨架；高频词 1；每天开口 5 句", "音标跟读录音"),
    ("P1 地基重建", "语法骨架：时态与句子结构；高频词 2", "能默写 12 个时态例句"),
    ("P1 地基重建", "精听入门：逐句听写；词汇 3", "听写正确率 ≥70%"),
    ("P1 地基重建", "口语启动：跟读 + 复述，克服开口恐惧", "阶梯阶段达标"),
    ("P1 地基重建", "学术高频词；阅读入门", "精读 5 篇"),
    ("P1 地基重建", "写作启动：50 → 100 词日记与段落", "累计 7 篇短文"),
    ("P1 地基重建", "考试题型全览 + 首次计时分项体验", "知道四项失分分布"),
    ("P1 地基重建", "阶段测评 1 + 全真模考", "拿到 baseline 分数"),
    ("P2 加速输出", "错题复盘；影子跟读法建立", "掌握 shadowing 流程"),
    ("P2 加速输出", "词汇扩张：词根词缀 + 搭配", "词汇量测试达标"),
    ("P2 加速输出", "精听升级：整篇演讲听写", "3 分钟材料听写 ≥80%"),
    ("P2 加速输出", "口语：2 分钟话题卡即兴", "7 段录音"),
    ("P2 加速输出", "阅读提速：外刊精读 3 篇 / 周", "阅读 180 wpm"),
    ("P2 加速输出", "写作：观点段与论证结构", "3 篇观点段"),
    ("P2 加速输出", "听力专项：长对话与讲座", "单项 ≥70%"),
    ("P2 加速输出", "阶段测评 2 + 阅读专项", "无字幕播客 75%"),
    ("P2 加速输出", "翻译专项：中国特色表达", "20 个高频译法"),
    ("P2 加速输出", "口语：连读弱读与语调（集中突破）", "跟读同步率提升"),
    ("P2 加速输出", "写作：作文模板个性化", "3 篇达目标分"),
    ("P2 加速输出", "半程大复盘；词汇冲刺", "半程报告"),
    ("P3 考试冲刺", "真题一刷（每周 1 套，计时）", "首套达线"),
    ("P3 考试冲刺", "真题一刷（第 2 套）+ 错题分类", "建立错题本"),
    ("P3 考试冲刺", "听力精练：错题精听 + 跟读", "听力 ≥75%"),
    ("P3 考试冲刺", "阅读精练：定位与同义替换", "阅读 ≥80%"),
    ("P3 考试冲刺", "写作冲刺：每周 3 篇 + 批改", "作文达目标分"),
    ("P3 考试冲刺", "阶段测评 3 + 翻译强化", "总分达线"),
    ("P3 考试冲刺", "真题二刷 + 每周计时模考", "稳定发挥"),
    ("P3 考试冲刺", "弱项专项（按错题分布决定）", "弱项正确率 +15%"),
    ("P3 考试冲刺", "全真模考周 1（严格计时）", "达标"),
    ("P3 考试冲刺", "全真模考周 2 + 作文面批", "达标"),
    ("P3 考试冲刺", "全真模考周 3 + 听力满分策略", "听力 ≥85%"),
    ("P3 考试冲刺", "全真模考周 4 + 阅读满速训练", "阅读 ≥85%"),
    ("P3 考试冲刺", "查漏补缺 + 高频词与模板回顾", "近 5 套均值达标"),
    ("P3 考试冲刺", "考前一周：模板固化、作息调整", "完考 + 考后复盘"),
    ("P4 母语化精修", "考后重启：全英输入日", "恢复节奏"),
    ("P4 母语化精修", "原版书精读 + 读书笔记", "读完半本"),
    ("P4 母语化精修", "口语：辩论与观点输出", "10 分钟长表达"),
    ("P4 母语化精修", "写作精修：地道语块替换中式表达", "1 篇 500 词"),
    ("P4 母语化精修", "阶段测评 4：高级样题", "达到目标等级边缘"),
    ("P4 母语化精修", "播客 / 纪录片沉浸 + 笔记法", "无字幕 85%"),
    ("P4 母语化精修", "口语：3 分钟即兴演讲周", "7 段演讲"),
    ("P4 母语化精修", "学术写作：essay 结构与引用", "1 篇 600 词"),
    ("P4 母语化精修", "听力极限：1.25–1.5 倍速训练", "倍速听懂 75%"),
    ("P4 母语化精修", "口音与韵律：台词 shadowing", "录音对比入库"),
    ("P4 母语化精修", "词汇升级：主动词汇检验 + 搭配", "测试通过"),
    ("P4 母语化精修", "阅读：整刊精读周", "一周读完一刊"),
    ("P4 母语化精修", "写作：长文逻辑打磨 + 改稿 3 轮", "1 篇精修长文"),
    ("P4 母语化精修", "专项突破（雅思 / 托福 / 职场英语任选）", "一项达标"),
    ("P4 母语化精修", "模拟目标等级测评 + 面试 / 会议场景", "场景自如"),
    ("P4 母语化精修", "全年输入输出总盘点", "数据总账"),
    ("P4 母语化精修", "录音对比与差距分析", "首月 vs 结业对比"),
    ("P4 母语化精修", "结业周：下一年计划 + 庆祝", "结业报告"),
]

WD = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
PHASE_FILL = {"W0": "FFF2CC", "P1": "EAF3FB", "P2": "EAF7EE", "P3": "FDF0E3", "P4": "F3EAFB"}


def parse_date(s: str) -> dt.date:
    return dt.datetime.strptime(s, "%Y-%m-%d").date()


def style_header(sheet, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = sheet.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def build(args) -> str:
    start = parse_date(args.start)
    weeks = args.weeks
    themes = (THEMES + THEMES)[:weeks]
    prep = args.prep_days
    prep_start = start - dt.timedelta(days=prep) if prep else None

    wb = Workbook()

    # ---------- 打卡表 ----------
    s = wb.active
    s.title = "打卡表"
    s.append(["日期", "星期", "周次", "阶段", "词汇(分)", "精听(分)", "阅读(分)", "口语(分)", "写作(分)",
              "合计(分)", "今日目标(分)", "完成率", "打卡", "备注"])
    rows = []
    for n in range(prep):
        rows.append((prep_start + dt.timedelta(days=n), 0, "W0 预备期"))
    for n in range(weeks * 7):
        rows.append((start + dt.timedelta(days=n), n // 7 + 1, themes[n // 7][0]))

    for i, (d, wk, phase) in enumerate(rows):
        r = i + 2
        s.append([d, WD[d.weekday()], wk, phase, None, None, None, None, None,
                  f"=SUM(E{r}:I{r})",
                  f'=IF(WEEKDAY(A{r},2)=7,{args.sun_min},IF(WEEKDAY(A{r},2)=6,{args.sat_min},{args.daily_min}))',
                  f'=IFERROR(J{r}/K{r},"")',
                  f'=IF(J{r}=0,"",IF(J{r}>=K{r}*0.8,"达标",IF(J{r}>=45,"部分","未完成")))',
                  None])
    for r in range(2, len(rows) + 2):
        s.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        s.cell(row=r, column=12).number_format = "0%"
        for c in range(1, 15):
            s.cell(row=r, column=c).border = BORDER
        for c in (2, 3, 13):
            s.cell(row=r, column=c).alignment = CENTER
        s.cell(row=r, column=4).fill = PatternFill("solid", fgColor=PHASE_FILL[s.cell(row=r, column=4).value[:2]])
    for col, w in zip("ABCDEFGHIJKLMN", [12, 7, 7, 14, 10, 10, 10, 10, 10, 10, 12, 9, 9, 46]):
        s.column_dimensions[col].width = w
    style_header(s, 14)

    # ---------- 周汇总 ----------
    w2 = wb.create_sheet("周汇总")
    w2.append(["周次", "起始日", "结束日", "阶段", "计划时长(分)", "实际时长(分)", "完成率", "打卡天数(≥45分)", "备注"])
    r = 2
    if prep:
        w2.append([0, prep_start, prep_start + dt.timedelta(days=prep - 1), "W0 预备期",
                   "=SUMIF(打卡表!$C:$C,$A2,打卡表!$K:$K)",
                   "=SUMIF(打卡表!$C:$C,$A2,打卡表!$J:$J)",
                   '=IFERROR(F2/E2,"")',
                   '=COUNTIFS(打卡表!$C:$C,$A2,打卡表!$J:$J,">=45")',
                   "装备期，不计入正式进度"])
        r = 3
    for i in range(1, weeks + 1):
        ws_ = start + dt.timedelta(days=(i - 1) * 7)
        w2.append([i, ws_, ws_ + dt.timedelta(days=6), themes[i - 1][0],
                   f"=SUMIF(打卡表!$C:$C,$A{r},打卡表!$K:$K)",
                   f"=SUMIF(打卡表!$C:$C,$A{r},打卡表!$J:$J)",
                   f'=IFERROR(F{r}/E{r},"")',
                   f'=COUNTIFS(打卡表!$C:$C,$A{r},打卡表!$J:$J,">=45")',
                   None])
        r += 1
    for rr in range(2, r):
        w2.cell(row=rr, column=2).number_format = "yyyy-mm-dd"
        w2.cell(row=rr, column=3).number_format = "yyyy-mm-dd"
        w2.cell(row=rr, column=7).number_format = "0%"
        for c in range(1, 10):
            w2.cell(row=rr, column=c).border = BORDER
    if prep:
        w2.cell(row=2, column=4).fill = PatternFill("solid", fgColor=PHASE_FILL["W0"])
    for col, w in zip("ABCDEFGHI", [7, 12, 12, 14, 13, 13, 9, 16, 26]):
        w2.column_dimensions[col].width = w
    style_header(w2, 9)

    # ---------- 指标看板 ----------
    w3 = wb.create_sheet("指标看板")
    w3.append(["指标", "起点", "一年目标", "当前值", "自测方式", "频率", "最近更新"])
    kpi = [
        ("被动词汇量", "", "", None, "词汇量测试 / 官方词表", "每 8 周", None),
        ("主动词汇", "", "", None, "随机抽 50 词造句", "每 8 周", None),
        ("听力理解（无字幕）", "", "85%+", None, "无字幕听 3 分钟播客，写 5 句摘要", "每周", None),
        ("口语流利度", "", "3 分钟无卡壳", None, "录音 + 统计卡顿/填充词", "每周", None),
        ("语速 (wpm)", "", "130–160", None, "录 1 分钟朗读数字数", "每月", None),
        ("阅读速度 (wpm)", "", "", None, "计时读 1 篇真题", "每 2 周", None),
        ("写作", "", "", None, "真题作文按评分标准自评", "每周", None),
        (f"{args.exam_name} 总分", "", "", None, "全真模考（严格计时）", "冲刺期每周", None),
    ]
    for k in kpi:
        w3.append(list(k))
    for rr in range(2, len(kpi) + 2):
        for c in range(1, 8):
            w3.cell(row=rr, column=c).border = BORDER
    for col, w in zip("ABCDEFG", [21, 12, 14, 20, 34, 12, 13]):
        w3.column_dimensions[col].width = w
    style_header(w3, 7)

    # ---------- 资源清单 ----------
    w4 = wb.create_sheet("资源清单")
    w4.append(["类别", "材料 / 工具", "用途", "使用阶段", "状态"])
    res = [
        ("词汇", f"{args.app}（选定一本词书）", "每天：先清「待复习」→ 学 15–20 新词；>120 张不加新词", "全周期", "未开始"),
        ("词汇", "搭配词典（如牛津搭配 / 柯林斯）", "查词看搭配", "全周期", "未开始"),
        ("词汇", "错词本 + 每日记录里的生词解释", "反复记不住的词 + 各造 1 句", "全周期", "未开始"),
        ("听力", "BBC Learning English – 6 Minute English", "精听入门", "P1", "未开始"),
        ("听力", "TED / TED-Ed", "整篇听写与影子跟读", "P1–P2", "未开始"),
        ("听力", "英语播客（新闻 / 访谈类）", "泛听 + 精听进阶", "P2–P4", "未开始"),
        ("听力", "纪录片（英文字幕 → 关字幕）", "无字幕沉浸", "P2–P4", "未开始"),
        ("口语", "AI 语音对话（每天 10 分钟）", "低压力开口", "全周期", "未开始"),
        ("口语", "外教课（每周 1–2 次，30 分钟）", "真实交流纠错", "P1 后半–P4", "未开始"),
        ("口语", "YouGlish", "查真实发音与语境", "全周期", "未开始"),
        ("口语", "发音教材（如 American Accent Training）", "连读、弱读、语调", "P1–P2", "未开始"),
        ("阅读", "每日推送原创文章（250–300 词 + 6 语块）", "每天 30′ 阅读 + 3 个任务", "全周期", "未开始"),
        ("阅读", "外刊（Economist / Guardian 等）", "精读，摘语块", "P2–P4", "未开始"),
        ("阅读", "原版书（从易到难 3 本）", "整本阅读", "P2–P4", "未开始"),
        ("写作", "每日写作题 + 教练批改", "中式表达→地道替换 / 语法硬伤 / 3 个升级句", "全周期", "未开始"),
        ("语法", "English Grammar in Use（中级 → 高级）", "每周 2 单元并做题", "P1–P2", "未开始"),
        ("考试", "近 10 年真题（≥3 刷）", "冲刺核心", "P1–P3", "未开始"),
        ("考试", "专项突破材料（听力 / 阅读 / 翻译）", "弱项突破", "P2–P3", "未开始"),
        ("考试", "官方报名与考试信息", "报名与考试时间", "P3", "未开始"),
    ]
    for row in res:
        w4.append(list(row))
    for rr in range(2, len(res) + 2):
        for c in range(1, 6):
            w4.cell(row=rr, column=c).border = BORDER
    for col, w in zip("ABCDE", [8, 46, 44, 14, 10]):
        w4.column_dimensions[col].width = w
    style_header(w4, 5)

    wb.save(args.out)
    return args.out


def main() -> None:
    p = argparse.ArgumentParser(description="English up · 每日打卡表生成器")
    p.add_argument("--out", required=True, help="输出 xlsx 路径")
    p.add_argument("--start", required=True, help="W1 周一日期 YYYY-MM-DD")
    p.add_argument("--weeks", type=int, default=52, help="周数，默认 52")
    p.add_argument("--prep-days", type=int, default=0, help="预备期天数（紧接 --start 之前），默认 0")
    p.add_argument("--app", default="词汇 App", help="词汇工具名")
    p.add_argument("--exam-name", default="目标考试", help="考试名称，用于指标看板")
    p.add_argument("--daily-min", type=int, default=150, help="周一到周五每日目标分钟")
    p.add_argument("--sat-min", type=int, default=210, help="周六目标分钟")
    p.add_argument("--sun-min", type=int, default=120, help="周日目标分钟")
    a = p.parse_args()

    out = build(a)
    start = parse_date(a.start)
    print(f"已生成：{out}")
    print(f"周期：{start}（周一）→ {start + dt.timedelta(days=a.weeks * 7 - 1)}（周日）｜"
          f"预备期 {a.prep_days} 天 ｜ 数据行 {a.prep_days + a.weeks * 7}")
    print("工作表：打卡表 / 周汇总 / 指标看板 / 资源清单")
    print("校验建议：任选一个工作日，手算 合计=五项之和、完成率=合计/目标；"
          "周汇总的 SUMIF 应等于该周 7 天「今日目标」之和。")


if __name__ == "__main__":
    main()
